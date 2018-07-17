#!/usr/bin/env python

import sys
import os
import re
import docx
import openpyxl

illegal_chars = '[０-９ａ-ｚＡ-Ｚ　！”＃＄％＆｀（）＊＋：；＜＞＝？]'

def check_docx(filename):
    doc = docx.Document(filename)
    illegal_line = re.compile(r'{}'.format(illegal_chars))
    for doc_paagraphs in doc.paragraphs:
        match = illegal_line.search(doc_paagraphs.text)
        if match != None:
            print(args[1] + "," + re.sub(r'({})'.format(illegal_chars), r'[\1]', doc_paagraphs.text))

def check_xlsx(filename):
    illegal_line = re.compile(r'{}'.format(illegal_chars))
    xlsx = openpyxl.load_workbook(filename)
    for sheetname in xlsx.sheetnames:
        sheet = xlsx[sheetname]
        for rows in sheet.rows:
            line = ''.join([cell.value if cell.value is not None else '' for cell in rows])
            match = illegal_line.search(line)
            if match != None:
                print(args[1] + "," + sheetname + "," + re.sub(r'({})'.format(illegal_chars), r'[\1]', line))

if __name__ == '__main__':
    args = sys.argv
    base, ext = os.path.splitext(args[1])

    if ext == '.docx' or ext == '.doc':
        check_docx(args[1])
    elif ext == '.xlsx' or ext == '.xlsm' or ext == '.xls':
        check_xlsx(args[1])
